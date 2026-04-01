"""Convert `.docx` files into markdown that PageIndex already knows how to parse."""

from __future__ import annotations

import re
from pathlib import Path


def _normalize_inline_text(text: str) -> str:
    """Collapse multi-line paragraph fragments into one clean inline string."""
    return " ".join(part.strip() for part in text.splitlines() if part.strip()).strip()


def _escape_markdown_cell(text: str) -> str:
    """Escape table-cell characters that would otherwise break markdown tables."""
    return _normalize_inline_text(text).replace("|", r"\|")


def _heading_level(style_name: str) -> int | None:
    """Infer markdown heading level from a Word style name."""
    match = re.search(r"heading\s+(\d+)", style_name, flags=re.IGNORECASE)
    if match:
        return min(int(match.group(1)), 6)

    if style_name.strip().lower() == "title":
        return 1

    return None


def _list_prefix(style_name: str) -> str:
    """Translate common Word list styles into markdown list prefixes."""
    normalized = style_name.strip().lower()
    bullet_match = re.search(r"list bullet(?:\s+(\d+))?$", normalized)
    if bullet_match:
        level = int(bullet_match.group(1) or "1")
        return f"{'  ' * (level - 1)}- "

    number_match = re.search(r"list number(?:\s+(\d+))?$", normalized)
    if number_match:
        level = int(number_match.group(1) or "1")
        return f"{'  ' * (level - 1)}1. "

    return ""


def _paragraph_to_markdown(paragraph) -> str:
    """Render one Word paragraph as markdown based on its style."""
    text = paragraph.text.strip()
    if not text:
        return ""

    style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
    level = _heading_level(style_name)
    if level is not None:
        return f"{'#' * level} {_normalize_inline_text(text)}"

    prefix = _list_prefix(style_name)
    if prefix:
        return f"{prefix}{_normalize_inline_text(text)}"

    return text


def _table_to_markdown(table) -> str:
    """Render a Word table into a simple markdown table."""
    rows: list[list[str]] = []
    for row in table.rows:
        normalized_row = [_escape_markdown_cell(cell.text) for cell in row.cells]
        if any(cell for cell in normalized_row):
            rows.append(normalized_row)

    if not rows:
        return ""

    col_count = max(len(row) for row in rows)
    padded_rows = [row + [""] * (col_count - len(row)) for row in rows]
    header = padded_rows[0]
    separator = ["---"] * col_count
    data_rows = padded_rows[1:]

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in data_rows)
    return "\n".join(lines)


def docx_to_markdown(file_path: str) -> str:
    """Convert a DOCX document into markdown for the markdown PageIndex path.

    Junior note:
    We do this because PageIndex has clean markdown ingestion support already,
    so converting DOCX to markdown is simpler than writing a brand-new DOCX tree builder.
    """
    try:
        from docx import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - dependency install issue
        raise ImportError(
            "python-docx is required for DOCX ingestion. Install it with "
            "`pip install python-docx` or reinstall from requirements.txt."
        ) from exc

    document = DocxDocument(file_path)
    blocks: list[str] = []

    if hasattr(document, "iter_inner_content"):
        iterable = document.iter_inner_content()
    else:  # pragma: no cover - backward compatibility fallback
        iterable = list(document.paragraphs) + list(document.tables)

    for block in iterable:
        if isinstance(block, Paragraph):
            rendered = _paragraph_to_markdown(block)
        elif isinstance(block, Table):
            rendered = _table_to_markdown(block)
        else:
            rendered = ""

        if rendered:
            blocks.append(rendered)

    markdown = "\n\n".join(blocks).strip()
    if not markdown:
        raise ValueError(f"No extractable text was found in DOCX file: {Path(file_path)}")

    return markdown + "\n"
